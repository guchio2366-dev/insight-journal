"""Derive national pine timber presentation and sale figures from retained INEGI for18."""
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data-source/atlas/mexico/agriculture/forestry/raw/ca2022_for18.xlsx'
OUTPUT = ROOT / 'src/data/atlas/mexico/pine-flow-2022.json'
EXPECTED_SHA = 'ae0e72e8598482e604984daa22e446a553766f2ff13441872290d2a588071a71'
if sha256(SOURCE.read_bytes()).hexdigest() != EXPECTED_SHA:
    raise ValueError('INEGI for18 source hash changed')

sheet = load_workbook(SOURCE, read_only=True, data_only=True).active
row = [sheet.cell(12, column).value for column in range(1, 20)]
if row[0] != '00 NAL' or row[1] != 'Pino (Pinus spp)':
    raise ValueError('National pine row changed')

def value(column):
    item = row[column - 1]
    if not isinstance(item, (int, float)) or item < 0:
        raise ValueError(f'Invalid numeric value at column {column}')
    return Decimal(str(item))

labels = [
    ('sawing', '製材向け', 6, 7),
    ('roundwood', '丸太形態', 8, 9),
    ('posts', '杭向け', 10, 11),
    ('firewood', '薪向け', 12, 13),
    ('charcoal', '炭向け', 14, 15),
    ('cellulose', 'セルロース向け', 16, 17),
    ('other', 'その他', 18, 19),
]
categories = [dict(id=id, label=label, obtainedM3=float(value(obtained)), soldM3=float(value(sold)))
              for id, label, obtained, sold in labels]
obtained, sold = value(4), value(5)
existing = json.loads((ROOT / 'src/data/atlas/mexico/agriculture.json').read_text())
if obtained != Decimal(str(existing['national']['pineObtainedM3'])):
    raise ValueError('for18 national pine volume differs from the for15 map indicator')
if sum((Decimal(str(item['obtainedM3'])) for item in categories), Decimal(0)) != obtained:
    raise ValueError('Obtained categories do not sum to national total')
if sum((Decimal(str(item['soldM3'])) for item in categories), Decimal(0)) != sold:
    raise ValueError('Sold categories do not sum to national total')
if sold > obtained or any(item['soldM3'] > item['obtainedM3'] for item in categories):
    raise ValueError('Sold volume exceeds obtained volume')

data = dict(
    schemaVersion=1,
    source='INEGI, Censo Agropecuario 2022, ca2022_for18',
    sourceUrl='https://www.inegi.org.mx/contenidos/programas/ca/2022/tabulados/ca2022_for18.xlsx',
    sourceSha256=EXPECTED_SHA,
    period='2021年10月〜2022年9月',
    unit='m³',
    species='Pino (Pinus spp.)',
    obtainedM3=float(obtained),
    soldM3=float(sold),
    categories=categories,
)
OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
print(f'{OUTPUT.relative_to(ROOT)}: {len(categories)} categories, {obtained} m³ obtained')
