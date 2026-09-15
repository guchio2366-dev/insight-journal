"""Build the reviewed 2020 USRC county-winner map.

Run:
  python scripts/population/build-religion-dominant.py RAW_DIR

RAW_DIR must contain the two official source files used by the audit:
  - 2020_USRC_Group_Detail.xlsx
  - cb_2020_us_county_5m.zip

The raw USRC workbook is intentionally not copied into the repository. The
published derivative contains only county FIPS, the winning display category,
and the source group code. Requires openpyxl and pyshp.
"""
from __future__ import annotations

import hashlib
import io
import json
import pathlib
import sys
import zipfile
from collections import Counter

import shapefile
from openpyxl import load_workbook

RAW = pathlib.Path(sys.argv[1])
ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA_OUT = ROOT / "data/atlas/population-religion-dominant-reviewed.json"
GEO_OUT = ROOT / "data/atlas/population-religion-ct-2020.geo.json"

MAINLINE = {
    "United Methodist Church",
    "Evangelical Lutheran Church in America",
    "American Baptist Churches in the USA",
    "Reformed Church in America",
    "United Church of Christ",
    "Episcopal Church",
    "Presbyterian Church (U.S.A.)",
}
BLACK_PROTESTANT = {
    "African Methodist Episcopal Church",
    "National Missionary Baptist Convention, Inc.",
    "Christian Methodist Episcopal Church",
}
OTHER = {"Vajarayana Buddhist", "Orthodox Church in America", "Hindu Yoga and Meditation"}
OTHER_CONSERVATIVE = {
    "American Baptist Association",
    "Amish Groups, undifferentiated",
    "Assemblies of God",
    "Christian Churches and Churches of Christ",
    "Christian Reformed Church in North America",
    "Christian and Missionary Alliance",
    "Church of Christ, Mennonite",
    "Church of God (Anderson, Indiana)",
    "Church of God (Cleveland, Tennessee)",
    "Church of the Brethren",
    "Church of the Nazarene",
    "Churches of Christ",
    "Convention of Original Free Will Baptist Churches",
    "Lehrerleut Hutterite",
    "Lutheran Church--Missouri Synod",
    "Mennonite Church USA",
    "National Association of Free Will Baptists",
    "Reformed Zion Union Apostolic Church",
    "Schmiedeleut Hutterite Group 2",
    "Seventh-day Adventist Church",
    "Weaverland Mennonite Conference",
    "Wesleyan Church",
    "Wisconsin Evangelical Lutheran Synod",
}

CATEGORIES = [
    ("catholic", "カトリック"),
    ("southern_baptist", "南部バプテスト"),
    ("mainline_protestant", "主流派プロテスタント"),
    ("nondenominational", "無教派キリスト教会"),
    ("other_conservative_protestant", "その他の保守系プロテスタント"),
    ("latter_day_saints", "末日聖徒"),
    ("black_protestant", "黒人プロテスタント"),
    ("other", "その他の最多グループ"),
    ("unreported", "報告なし"),
]
EXPECTED_COUNTS = {
    "catholic": 1221,
    "southern_baptist": 1056,
    "mainline_protestant": 327,
    "nondenominational": 274,
    "other_conservative_protestant": 116,
    "latter_day_saints": 94,
    "black_protestant": 15,
    "other": 2,
    "unreported": 3,
}


def fips(value: object) -> str | None:
    try:
        return str(int(value)).zfill(5)
    except (TypeError, ValueError):
        return None


def category(name: str) -> str:
    if name == "Catholic Church":
        return "catholic"
    if name == "Southern Baptist Convention":
        return "southern_baptist"
    if name == "Non-denominational Christian Churches":
        return "nondenominational"
    if name == "Church of Jesus Christ of Latter-day Saints":
        return "latter_day_saints"
    if name in MAINLINE:
        return "mainline_protestant"
    if name in BLACK_PROTESTANT:
        return "black_protestant"
    if name in OTHER_CONSERVATIVE:
        return "other_conservative_protestant"
    if name in OTHER:
        return "other"
    raise ValueError(f"Unreviewed winning group: {name}")


def county_winners(path: pathlib.Path):
    sheet = load_workbook(path, read_only=True, data_only=True)["2020 Group by County"]
    best: dict[str, tuple[float, str, str]] = {}
    tied: set[str] = set()
    anomaly_seen = False
    for row in sheet.iter_rows(min_row=2, values_only=True):
        code = fips(row[0])
        if code is None:
            continue
        state, county, group_code, group_name, adherents = row[1], row[2], row[3], row[4], row[6]
        if code == "17032" and state == "Illinois" and county == "Cook County" and group_name == "Chabad Judaism":
            anomaly_seen = True
            continue
        if not isinstance(adherents, (int, float)) or adherents <= 0:
            continue
        candidate = (float(adherents), str(group_code), str(group_name))
        previous = best.get(code)
        if previous is None or candidate[0] > previous[0]:
            best[code] = candidate
            tied.discard(code)
        elif candidate[0] == previous[0]:
            tied.add(code)
    if not anomaly_seen:
        raise ValueError("Expected Cook County source anomaly was not found")
    if tied:
        raise ValueError(f"Tied county winners require review: {sorted(tied)}")
    return best


def shape_reader(path: pathlib.Path):
    archive = zipfile.ZipFile(path)
    stem = next(name[:-4] for name in archive.namelist() if name.endswith(".shp"))
    return shapefile.Reader(
        shp=io.BytesIO(archive.read(stem + ".shp")),
        shx=io.BytesIO(archive.read(stem + ".shx")),
        dbf=io.BytesIO(archive.read(stem + ".dbf")),
        encoding="utf-8",
    )


def rounded(value):
    if isinstance(value, dict):
        return {key: rounded(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [rounded(item) for item in value]
    return round(value, 4) if isinstance(value, float) else value


def main():
    workbook = RAW / "2020_USRC_Group_Detail.xlsx"
    boundary = RAW / "cb_2020_us_county_5m.zip"
    winners = county_winners(workbook)
    features = []
    rows = []
    geometry_ids = set()
    for item in shape_reader(boundary).iterShapeRecords():
        record = item.record.as_dict()
        state = record["STATEFP"]
        if state in {"02", "15"} or int(state) > 56:
            continue
        code = record["GEOID"]
        if code in geometry_ids:
            raise ValueError(f"Duplicate boundary FIPS: {code}")
        geometry_ids.add(code)
        winner = winners.get(code)
        display = "unreported" if winner is None else category(winner[2])
        rows.append({"id": "county:" + code, "category": display, "group": None if winner is None else winner[1]})
        features.append({
            "type": "Feature",
            "properties": {"id": "county:" + code, "name": record["NAMELSAD"]},
            "geometry": rounded(item.shape.__geo_interface__),
        })
    if len(geometry_ids) != 3108:
        raise ValueError(f"Expected 3,108 contiguous county equivalents, got {len(geometry_ids)}")
    counts = Counter(row["category"] for row in rows)
    if dict(counts) != EXPECTED_COUNTS:
        raise ValueError(f"Reviewed category totals changed: {dict(counts)}")
    if {row["id"][7:] for row in rows} != geometry_ids:
        raise ValueError("Religion data and boundary FIPS do not match")

    source_hashes = {
        path.name: {"bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        for path in (workbook, boundary)
    }
    output = {
        "version": 1,
        "year": 2020,
        "universe": "congregation-linked adherents",
        "geography": "2020 county equivalents; contiguous 48 states and District of Columbia",
        "source": {
            "url": "https://www.usreligioncensus.org/node/1639",
            "title": "2020 U.S. Religion Census",
            "citation": "Grammich et al. 2023, 2020 U.S. Religion Census",
        },
        "categories": [
            {"id": category_id, "label": label, "count": EXPECTED_COUNTS[category_id]}
            for category_id, label in CATEGORIES
        ],
        "rows": sorted(rows, key=lambda row: row["id"]),
        "audit": {"sourceFiles": source_hashes, "excludedAnomaly": "Illinois Cook County Chabad row with FIPS 17032"},
    }
    DATA_OUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n")
    geo = {"type": "FeatureCollection", "features": sorted(features, key=lambda feature: feature["properties"]["id"])}
    # The shared map already has every 2020 county except Connecticut's former
    # eight counties. Publish only that small override; prepare-religion.mjs
    # replaces the nine current planning regions deterministically at build time.
    patch = {"type": "FeatureCollection", "features": [feature for feature in geo["features"] if feature["properties"]["id"].startswith("county:09")]}
    if len(patch["features"]) != 8:
        raise ValueError(f"Expected eight 2020 Connecticut counties, got {len(patch['features'])}")
    GEO_OUT.write_text(json.dumps(patch, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(json.dumps({"rows": len(rows), "counts": counts, "connecticutFeatures": len(patch["features"]), "patchBytes": GEO_OUT.stat().st_size}, default=dict))


if __name__ == "__main__":
    main()
