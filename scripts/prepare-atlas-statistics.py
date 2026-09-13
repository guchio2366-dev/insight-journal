#!/usr/bin/env python3
"""Normalize reviewed USDA files for the North America agriculture page.

Large Farm Income and PSD inputs are downloaded when updating and are not
committed. Small FATUS workbooks and the reviewed GATS rice extract live under
data-source/atlas. This script never fetches data and never fills missing rows
with zero.
"""
import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--farm-income", required=True, type=Path)
    parser.add_argument("--psd", required=True, type=Path)
    parser.add_argument("--calendar", type=Path, default=Path("data-source/atlas/ers-fatus-calendar-2025.xlsx"))
    parser.add_argument("--markets", type=Path, default=Path("data-source/atlas/ers-fatus-top-markets-2026-09.xlsx"))
    parser.add_argument("--rice", type=Path, default=Path("data-source/atlas/rice-gats-2025.csv"))
    parser.add_argument("--generated-at", default="2026-09-13")
    parser.add_argument("--output", type=Path, default=Path("src/data/atlas/statistics-generated.ts"))
    return parser.parse_args()


def cash_receipts(path):
    keys = {"all":"CRAUSAC--VAP","crops":"CRAUSCO--VAP","livestock":"CRAUSLV--VAP","corn":"CRAUSCR--VAP","soybean":"CRAUSSY--VAP","fruit-nuts":"CRAUSFN--VAP","vegetables":"CRAUSVG--VAP","food-grains":"CRAUSFO--VAP","feed":"CRAUSFE--VAP","floriculture":"CRAUSGNFCVAP","cotton":"CRAUSCT--VAP","sugar-cane":"CRAUSCW--VAP","sugar-beets":"CRAUSSR--VAP","other-oil":"CRAUSOC--VAP","other-crops":"CRAUSAOCOVAP","misc-crops":"CRAUSAOOTVAP","tobacco":"CRAUSTB--VAP","cattle":"CRAUSCL--VAP","poultry":"CRAUSPG--VAP","milk":"CRAUSDY--VAP","hogs":"CRAUSHG--VAP","other-animals":"CRAUSLVMIVAP"}
    selected = {}
    with path.open(encoding="cp1252", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["Year"] != "2025" or row["State"] != "US":
                continue
            if row["artificialKey"] in keys.values():
                if row["artificialKey"] in selected: raise ValueError(f"duplicate Farm Income key: {row['artificialKey']}")
                if not row["unit_desc"].strip().startswith("$1,000"): raise ValueError(f"unexpected Farm Income unit: {row['unit_desc']}")
                selected[row["artificialKey"]] = int(row["Amount"])
    missing=set(keys.values())-set(selected)
    if missing: raise ValueError(f"Farm Income input lacks required keys: {sorted(missing)}")
    value=lambda name:selected[keys[name]]
    if value("all") != value("crops") + value("livestock"): raise ValueError("parent totals do not reconcile")
    crop_items=[("corn","とうもろこし",value("corn")),("soybean","大豆",value("soybean")),("fruit-nuts","果物・ナッツ",value("fruit-nuts")),("vegetables","野菜・メロン",value("vegetables")),("food-grains","小麦・米など",value("food-grains")),("other-feed","牧草等",value("feed")-value("corn")),("floriculture","花き",value("floriculture")),("cotton","綿花",value("cotton")),("sugar-crops","砂糖原料",value("sugar-cane")+value("sugar-beets")),("other-published","落花生・菜種等",value("other-oil")-value("soybean")+value("other-crops")-value("misc-crops")-value("floriculture")-value("sugar-cane")-value("sugar-beets")+value("tobacco")),("miscellaneous","その他",value("misc-crops"))]
    animal_items=[("cattle-calves","牛・子牛",value("cattle")),("poultry-eggs","家禽・卵",value("poultry")),("milk","生乳",value("milk")),("hogs","豚",value("hogs")),("other-animals","その他",value("other-animals"))]
    make=lambda items:[{"id":i,"label":label,"valueThousandUsd":amount} for i,label,amount in items]
    groups=[{"id":"crops","label":"作物","totalThousandUsd":value("crops"),"items":make(crop_items),"reconciliationThousandUsd":value("crops")-sum(x[2] for x in crop_items)},{"id":"livestock","label":"畜産","totalThousandUsd":value("livestock"),"items":make(animal_items),"reconciliationThousandUsd":value("livestock")-sum(x[2] for x in animal_items)}]
    for group in groups:
        if abs(group["reconciliationThousandUsd"]) > len(group["items"]): raise ValueError(f"{group['id']} children exceed rounding tolerance")
    return {"year":2025,"status":"estimate","unit":"thousand USD","totalThousandUsd":value("all"),"groups":groups}


def calendar_trade(path):
    sheet = load_workbook(path, data_only=True, read_only=True).active
    series = []
    for row in sheet.iter_rows(values_only=True):
        if isinstance(row[0], int) and 2016 <= row[0] <= 2025:
            series.append({"year": row[0], "exports": round(row[1] / 1000, 3), "imports": round(row[2] / 1000, 3)})
    if [row["year"] for row in series] != list(range(2016, 2026)):
        raise ValueError("FATUS calendar input must contain every year from 2016 through 2025")
    return {"period": "calendar year", "unit": "billion USD", "series": series}


def export_markets(path, rice_path):
    sheet = load_workbook(path, data_only=True, read_only=True)["Feb. update of Dec. 2025 data"]
    mapping = {"Soybeans": "soybean", "Corn": "corn", "Wheat, unmilled": "wheat", "Cotton, excluding linters": "cotton"}
    result, current = {}, None
    for row in sheet.iter_rows(values_only=True):
        label = row[0]
        if label in mapping:
            current = mapping[label]
            result[current] = {"basis": label, "year": 2025, "unit": "metric tons", "rows": [], "total": None}
        elif current and label == "World total":
            result[current]["total"] = float(row[2])
            current = None
        elif current and isinstance(row[2], (int, float)):
            result[current]["rows"].append({"country": str(label), "value": float(row[2])})
    for crop, record in result.items():
        if record["total"] is None or len(record["rows"]) < 5:
            raise ValueError(f"FATUS markets input lacks required {crop} rows")
        top = record.pop("rows")[:5]
        record["destinations"] = top + [{"country": "Other", "value": round(record["total"] - sum(row["value"] for row in top), 1)}]

    with rice_path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    total_rows = [row for row in rows if row["country"] == "World total"]
    destinations = [row for row in rows if row["country"] != "World total"]
    if len(total_rows) != 1 or len(destinations) != 6:
        raise ValueError("GATS rice extract must contain six destination rows and one world total")
    total = float(total_rows[0]["value_metric_tons_milled_equivalent"])
    parsed = [{"country": row["country"], "value": float(row["value_metric_tons_milled_equivalent"])} for row in destinations]
    if abs(sum(row["value"] for row in parsed) - total) > 0.11:
        raise ValueError("GATS rice destinations do not sum to the reviewed total")
    result["rice"] = {
        "basis": "Rice, BICO-HS10, FAS-converted line items",
        "year": 2025,
        "unit": "metric tons, milled-equivalent (MTMEQ)",
        "total": total,
        "destinations": parsed,
    }
    return result


def production(path):
    commodities = {"corn": "Corn", "soybean": "Oilseed, Soybean", "wheat": "Wheat", "cotton": "Cotton", "rice": "Rice, Milled"}
    records = {crop: {year: [] for year in range(2015, 2025)} for crop in commodities}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["Attribute_Description"] != "Production":
                continue
            crop = next((crop for crop, name in commodities.items() if name == row["Commodity_Description"]), None)
            year = int(row["Market_Year"])
            if crop and year in records[crop]:
                records[crop][year].append({"code": row["Country_Code"], "country": row["Country_Name"], "value": float(row["Value"]), "unit": row["Unit_Description"]})
    output = {}
    for crop, years in records.items():
        trend = []
        for year, rows in years.items():
            if not rows:
                raise ValueError(f"PSD input lacks {crop} production for marketing year {year}")
            world = sum(row["value"] for row in rows)
            us_rows = [row for row in rows if row["code"] == "US"]
            if len(us_rows) != 1:
                raise ValueError(f"PSD input must have exactly one US {crop} row for {year}")
            us = us_rows[0]["value"]
            trend.append({"year": year, "us": us, "world": world, "share": round(us / world * 100, 3)})
        latest = years[2024]
        world = sum(row["value"] for row in latest)
        countries = sorted(latest, key=lambda row: row["value"], reverse=True)[:5]
        if not any(row["code"] == "US" for row in countries):
            countries.append(next(row for row in latest if row["code"] == "US"))
        output[crop] = {
            "basis": commodities[crop],
            "unit": latest[0]["unit"],
            "marketYear": 2024,
            "worldTotal": world,
            "countries": [{"code": row["code"], "country": row["country"], "value": row["value"]} for row in countries],
            "trend": trend,
        }
    return output


def main():
    args = parse_args()
    result = {
        "generatedAt": args.generated_at,
        "national": {"cashReceipts": cash_receipts(args.farm_income), "trade": calendar_trade(args.calendar)},
        "exports": export_markets(args.markets, args.rice),
        "production": production(args.psd),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("export const atlasStatistics = " + json.dumps(result, ensure_ascii=False, indent=2) + " as const;\n", encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
